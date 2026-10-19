import argparse
import os
from collections import Counter
from pathlib import Path
from urllib.parse import quote_plus

from bson import ObjectId
from openpyxl import load_workbook
from pymongo import MongoClient, UpdateOne


SOURCE = "ReviewClean"
BACKEND_ENV = Path(__file__).resolve().parents[1] / "backend-node" / ".env"


def normalize_id(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def get_mongodb_uri():
    uri = os.getenv("MONGODB_URI")
    if uri:
        return uri

    try:
        from dotenv import load_dotenv
    except ImportError as error:
        raise RuntimeError(
            "Set MONGODB_URI or install python-dotenv to read backend-node/.env."
        ) from error

    load_dotenv(BACKEND_ENV)
    uri = os.getenv("MONGODB_URI")
    if uri:
        return uri

    password = os.getenv("MONGO_DB")
    if not password:
        raise RuntimeError("Set MONGODB_URI or MONGO_DB in backend-node/.env.")

    encoded_password = quote_plus(password)
    return (
        "mongodb+srv://novashopdb:"
        f"{encoded_password}@novashopdb.1rntkxm.mongodb.net/"
        "novashopdb?appName=NovaShopDB"
    )


def read_review_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["Reviews"]
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("The Reviews sheet is empty.")

        columns = {str(value).strip(): index for index, value in enumerate(header) if value}
        required = {"customer_id", "product_id", "rating"}
        missing = required - columns.keys()
        if missing:
            raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")

        for row in rows:
            customer_id = normalize_id(row[columns["customer_id"]])
            item_id = normalize_id(row[columns["product_id"]])
            rating_value = row[columns["rating"]]
            if not customer_id or not item_id or rating_value is None:
                yield None
                continue
            rating = int(rating_value) if float(rating_value).is_integer() else None
            if rating is None or rating < 1 or rating > 5:
                yield None
                continue
            yield customer_id, item_id, rating
    finally:
        workbook.close()


def normalize_item_id_query(item_id):
    values = [item_id]
    if item_id.isdigit():
        values.append(int(item_id))
    return values


def main():
    parser = argparse.ArgumentParser(
        description="Import source ratings into datasetratings without creating user accounts."
    )
    parser.add_argument("--reviews", required=True, type=Path, help="Path to ReviewClean.xlsx")
    parser.add_argument(
        "--products-json",
        required=True,
        type=Path,
        help="Path to the catalog products.json",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Write mapped ratings to MongoDB. Without this flag, only report the mapping.",
    )
    args = parser.parse_args()

    import json

    with args.products_json.open(encoding="utf-8") as catalog_file:
        catalog = json.load(catalog_file)
    catalog_item_ids = {
        normalize_id(product.get("item_id"))
        for product in catalog
        if isinstance(product, dict) and product.get("item_id") is not None
    }

    row_count = 0
    invalid_rows = 0
    customer_ids = set()
    review_counts_by_item = Counter()
    pairs = set()
    duplicate_pairs = 0
    for row in read_review_rows(args.reviews):
        row_count += 1
        if row is None:
            invalid_rows += 1
            continue
        customer_id, item_id, _ = row
        pair = (customer_id, item_id)
        if pair in pairs:
            duplicate_pairs += 1
        pairs.add(pair)
        customer_ids.add(customer_id)
        review_counts_by_item[item_id] += 1

    if duplicate_pairs:
        raise ValueError(
            f"Found {duplicate_pairs} duplicate customer/product pairs; "
            "resolve them before importing."
        )

    mongo_client = MongoClient(get_mongodb_uri(), serverSelectionTimeoutMS=10000)
    try:
        mongo_client.admin.command("ping")
        database = mongo_client["novashopdb"]
        products_collection = database["products"]
        ratings_collection = database["datasetratings"]

        item_id_values = [
            value
            for item_id in review_counts_by_item
            for value in normalize_item_id_query(item_id)
        ]
        db_product_matches = {}
        for product in products_collection.find(
            {"item_id": {"$in": item_id_values}},
            {"_id": 1, "item_id": 1},
        ):
            item_id = normalize_id(product.get("item_id"))
            db_product_matches.setdefault(item_id, []).append(product["_id"])

        unique_db_products = {
            item_id: matches[0]
            for item_id, matches in db_product_matches.items()
            if len(matches) == 1
        }
        duplicate_db_item_ids = {
            item_id for item_id, matches in db_product_matches.items() if len(matches) > 1
        }
        catalog_matched_ids = set(review_counts_by_item) & catalog_item_ids
        db_matched_ids = set(review_counts_by_item) & set(unique_db_products)
        importable_ids = catalog_matched_ids & db_matched_ids

        print(f"Review rows: {row_count}")
        print(f"Invalid rows: {invalid_rows}")
        print(f"Distinct source customers: {len(customer_ids)}")
        print(f"Distinct reviewed product IDs: {len(review_counts_by_item)}")
        print(f"Product IDs found in products.json: {len(catalog_matched_ids)}")
        print(f"Product IDs found exactly once in MongoDB: {len(db_matched_ids)}")
        print(f"Product IDs with duplicate MongoDB item_id: {len(duplicate_db_item_ids)}")
        print(
            "Importable ratings: "
            f"{sum(review_counts_by_item[item_id] for item_id in importable_ids)}"
        )
        print(
            "Unmatched catalog IDs: "
            f"{len(set(review_counts_by_item) - catalog_item_ids)}"
        )
        print(
            "IDs missing or ambiguous in MongoDB: "
            f"{len(set(review_counts_by_item) - db_matched_ids)}"
        )
        if not args.apply:
            print("Dry run only; no database documents were written. Add --apply to import.")
            return

        ratings_collection.create_index(
            [("source", 1), ("customerId", 1), ("productItemId", 1)],
            unique=True,
            name="source_customer_product_unique",
        )
        ratings_collection.create_index(
            [("productId", 1), ("source", 1)],
            name="product_source_lookup",
        )

        operations = []
        inserted_total = 0
        updated_total = 0

        def flush_batch():
            nonlocal inserted_total, updated_total
            if not operations:
                return
            result = ratings_collection.bulk_write(operations, ordered=False)
            inserted_total += result.upserted_count
            updated_total += result.modified_count
            operations.clear()

        for row in read_review_rows(args.reviews):
            if row is None:
                continue
            customer_id, item_id, rating = row
            product_id = unique_db_products.get(item_id)
            if item_id not in importable_ids or product_id is None:
                continue
            operations.append(
                UpdateOne(
                    {
                        "source": SOURCE,
                        "customerId": customer_id,
                        "productItemId": item_id,
                    },
                    {
                        "$set": {
                            "productId": ObjectId(product_id),
                            "rating": rating,
                        },
                        "$setOnInsert": {
                            "source": SOURCE,
                            "customerId": customer_id,
                            "productItemId": item_id,
                        },
                    },
                    upsert=True,
                )
            )
            if len(operations) >= 1000:
                flush_batch()

        flush_batch()
        print(f"Import results: {inserted_total} inserted, {updated_total} updated.")
        print("Import complete. Dataset customer IDs were not added to the users collection.")
    finally:
        mongo_client.close()


if __name__ == "__main__":
    main()
